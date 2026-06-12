#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
MVP датасета 3D-моделей ТС: топ-1600 уникальных авто на уровне
mark + model + generation, дающий ≥90% покрытия объявлений на дорогах.

Источник: priority_export.csv (9 741 строк, развёрнутые по типам кузова и
рестайлингам). Свёртка → ~2 500 уникальных авто, дальше отбор топ-1 600.
"""
from pathlib import Path

import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

SRC = Path('/sessions/jolly-charming-pasteur/mnt/cars_analysis/priority_export.csv')
OUT = Path('/sessions/jolly-charming-pasteur/mnt/cars_analysis/MVP_top1600_3D_dataset.xlsx')
TOP_N = 1600

BRAND_DISPLAY = {
    'lada': 'LADA', 'vaz': 'ВАЗ', 'uaz': 'УАЗ', 'gaz': 'ГАЗ', 'kamaz': 'КамАЗ',
    'mercedes benz': 'Mercedes-Benz', 'mercedes-benz': 'Mercedes-Benz',
    'aston martin': 'Aston Martin', 'aston_martin': 'Aston Martin',
    'land rover': 'Land Rover', 'land_rover': 'Land Rover',
    'alfa romeo': 'Alfa Romeo', 'rolls-royce': 'Rolls-Royce',
    'great_wall': 'Great Wall', 'great wall': 'Great Wall',
    'ssangyong': 'SsangYong', 'bmw': 'BMW', 'kia': 'Kia',
}


def display_brand(b):
    if not isinstance(b, str):
        return b
    key = b.strip().lower()
    return BRAND_DISPLAY.get(key, key.replace('_', ' ').title())


def main():
    df = pd.read_csv(SRC, encoding='utf-8-sig')
    print(f'Исходник: {len(df):,} строк')

    # ============================================================
    # СВЁРТКА: 1 строка = марка + модель + поколение
    # ============================================================
    grp = df.groupby(['mark', 'model', 'generation'], sort=False)
    agg = grp.agg(
        restyling_min=('restyling', 'min'),
        restyling_max=('restyling', 'max'),
        body_types_set=('body_type', lambda s: sorted(set(s.dropna().astype(str)))),
        year_from=('year_from', 'min'),
        year_to=('year_to', 'max'),
        origin_category=('origin_category', 'first'),
        vehicle_type_category=('vehicle_type_category', 'first'),
        ads_count=('ads_count', 'max'),        # значение одинаковое внутри группы
        avg_price=('avg_price_inflation_adj', 'max'),
        avg_age=('avg_age', 'mean'),
        priority_score=('priority_score', 'max'),
        has_real_metrics=('has_real_metrics', 'any'),
        score_origin=('score_origin', 'max'),
        score_type=('score_type', 'max'),
        score_ads=('score_ads', 'max'),
        score_price=('score_price', 'max'),
        score_age=('score_age', 'max'),
        n_variants=('body_type', 'size'),
    ).reset_index()

    # представления
    agg['restylings'] = agg.apply(
        lambda r: (str(int(r['restyling_min']))
                   if r['restyling_min'] == r['restyling_max']
                   else f'{int(r["restyling_min"])}–{int(r["restyling_max"])}'),
        axis=1
    )
    agg['body_types'] = agg['body_types_set'].apply(lambda s: ', '.join(s))
    agg['mark_display'] = agg['mark'].apply(display_brand)
    agg['vehicle_id'] = (
        agg['mark'].str.replace(' ', '_').str.lower()
        + '__' + agg['model'].str.replace(' ', '_').str.lower()
        + '__g' + agg['generation'].astype(int).astype(str)
    )

    print(f'После свёртки: {len(agg):,} уникальных авто (mark+model+generation)')

    # ============================================================
    # СОРТИРОВКА И ОТБОР ТОП-1600
    # ============================================================
    agg = agg.sort_values(
        ['ads_count', 'priority_score'], ascending=[False, False]
    ).reset_index(drop=True)

    total_ads = agg['ads_count'].sum()
    agg['coverage_%'] = agg['ads_count'].cumsum() / total_ads * 100
    agg['rank'] = range(1, len(agg) + 1)

    top = agg.head(TOP_N).copy()
    cov_top = top['coverage_%'].iloc[-1]
    print(f'Топ-{TOP_N}: покрытие {cov_top:.2f}% по ads_count')
    if cov_top < 90:
        # подцепляем хвост до 90%
        need = (agg['coverage_%'] <= 90).sum() + 1
        print(f'   ⚠ < 90% — расширяю до топ-{need}')
        top = agg.head(need).copy()

    # ============================================================
    # ПЛЕЙСХОЛДЕРЫ ДЛЯ СЛЕДУЮЩИХ ЭТАПОВ
    # ============================================================
    placeholders = {
        # Этап 3 — габариты
        'length_mm': None, 'width_mm': None, 'height_mm': None,
        'wheelbase_mm': None, 'ground_clearance_mm': None,
        'dimensions_source': None,
        # Этап 4 — 3D-модель
        'model_3d_status': 'не начат',
        'model_3d_source': None,        # sketchfab/hum3d/meshy/...
        'model_3d_license': None,
        'model_file_path': None,
        'model_format': None,           # blend/fbx/obj/glb
        # Этап 5 — валидация в Blender
        'scale_check': None,            # ok / off / нет данных
        'blender_compatible': None,
        # Заметки
        'note': '',
    }
    for col, default in placeholders.items():
        top[col] = default

    # ============================================================
    # СБОРКА ИТОГОВОГО ДАТАФРЕЙМА
    # ============================================================
    out = top[[
        'rank', 'priority_score', 'coverage_%', 'vehicle_id',
        'mark_display', 'model', 'generation', 'restylings', 'body_types',
        'year_from', 'year_to', 'n_variants',
        'origin_category', 'vehicle_type_category',
        'ads_count', 'avg_price', 'avg_age', 'has_real_metrics',
        'score_origin', 'score_type', 'score_ads', 'score_price', 'score_age',
        'length_mm', 'width_mm', 'height_mm',
        'wheelbase_mm', 'ground_clearance_mm', 'dimensions_source',
        'model_3d_status', 'model_3d_source', 'model_3d_license',
        'model_file_path', 'model_format',
        'scale_check', 'blender_compatible',
        'note',
    ]].copy()

    out = out.rename(columns={
        'rank': 'Ранг',
        'priority_score': 'Балл приоритета',
        'coverage_%': 'Покрытие, %',
        'vehicle_id': 'ID',
        'mark_display': 'Марка',
        'model': 'Модель',
        'generation': 'Поколение',
        'restylings': 'Рестайлинги',
        'body_types': 'Типы кузова',
        'year_from': 'Год от',
        'year_to': 'Год до',
        'n_variants': 'Вариантов кузовов',
        'origin_category': 'Локализация',
        'vehicle_type_category': 'Тип ТС',
        'ads_count': 'Кол-во объявлений',
        'avg_price': 'Средняя цена, ₽ (с инфл.)',
        'avg_age': 'Средний возраст, лет',
        'has_real_metrics': 'Реальные метрики',
        'score_origin': 'Б: локализация',
        'score_type': 'Б: тип ТС',
        'score_ads': 'Б: объявления',
        'score_price': 'Б: цена',
        'score_age': 'Б: возраст',
        'length_mm': 'Длина, мм',
        'width_mm': 'Ширина, мм',
        'height_mm': 'Высота, мм',
        'wheelbase_mm': 'Колёсная база, мм',
        'ground_clearance_mm': 'Клиренс, мм',
        'dimensions_source': 'Источник габаритов',
        'model_3d_status': 'Статус 3D',
        'model_3d_source': 'Источник 3D',
        'model_3d_license': 'Лицензия 3D',
        'model_file_path': 'Путь к файлу',
        'model_format': 'Формат',
        'scale_check': 'Проверка масштаба',
        'blender_compatible': 'Совместимо с Blender',
        'note': 'Комментарий',
    })

    # ============================================================
    # ГРУППЫ КОЛОНОК ДЛЯ XLSX
    # ============================================================
    GROUPS = [
        ('Приоритет',                'DCE6F1', ['Ранг', 'Балл приоритета', 'Покрытие, %', 'ID']),
        ('Авто',                     'E2EFDA', ['Марка', 'Модель', 'Поколение', 'Рестайлинги',
                                                 'Типы кузова', 'Год от', 'Год до', 'Вариантов кузовов']),
        ('Категории',                'FFF2CC', ['Локализация', 'Тип ТС']),
        ('Метрики',                  'FCE4D6', ['Кол-во объявлений', 'Средняя цена, ₽ (с инфл.)',
                                                 'Средний возраст, лет', 'Реальные метрики']),
        ('Компоненты оценки (0–1)', 'EDEDED', ['Б: локализация', 'Б: тип ТС', 'Б: объявления',
                                                 'Б: цена', 'Б: возраст']),
        ('Габариты (заполнить)',     'D9E1F2', ['Длина, мм', 'Ширина, мм', 'Высота, мм',
                                                 'Колёсная база, мм', 'Клиренс, мм',
                                                 'Источник габаритов']),
        ('3D-модель (заполнить)',    'FFE699', ['Статус 3D', 'Источник 3D', 'Лицензия 3D',
                                                 'Путь к файлу', 'Формат']),
        ('Валидация (заполнить)',    'F8CBAD', ['Проверка масштаба', 'Совместимо с Blender']),
        ('Заметки',                  'F2F2F2', ['Комментарий']),
    ]

    # ============================================================
    # ЗАПИСЬ EXCEL: лист "MVP top-1600" + "Описание" + "Метрики"
    # ============================================================
    wb = Workbook()
    wb.remove(wb.active)

    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # ---- Лист "Описание" ----
    ws_info = wb.create_sheet('Описание')
    info_rows = [
        ['Проект', 'Датасет 3D-моделей транспортных средств'],
        ['Файл', OUT.name],
        ['', ''],
        ['Источник', 'priority_export.csv (9 741 строк, развёрнутые по кузовам)'],
        ['Свёртка', 'mark + model + generation → одна строка'],
        ['Уникальных авто', f'{len(agg):,}'.replace(',', ' ')],
        ['Отобрано в MVP', f'топ-{len(out)}'],
        ['Покрытие MVP', f'{out["Покрытие, %"].iloc[-1]:.2f}% объявлений на дорогах'],
        ['', ''],
        ['Группы колонок', ''],
        ['  Приоритет', 'Ранг, балл, накопленное покрытие %, ID'],
        ['  Авто', 'Марка, модель, поколение, рестайлинги, типы кузовов, годы'],
        ['  Категории', 'Локализация (Российские / Локализованные / Импорт) и тип ТС'],
        ['  Метрики', 'Кол-во объявлений, средняя цена ₽ (приведена к 2025), возраст'],
        ['  Компоненты оценки', '5 нормированных баллов 0–1, из которых собран приоритет'],
        ['  Габариты', 'ПУСТЫЕ — этап 3, длина/ширина/высота/база/клиренс'],
        ['  3D-модель', 'ПУСТЫЕ — этап 4, источник, файл, формат'],
        ['  Валидация', 'ПУСТЫЕ — этап 5, проверка масштаба в Blender'],
        ['  Заметки', 'Свободный комментарий'],
        ['', ''],
        ['Следующие этапы', ''],
        ['  Этап 3', 'Заполнить габариты (скрейп drom/wiki + ручная верификация)'],
        ['  Этап 4', 'Источники 3D: Sketchfab/BlenderKit → Hum3D/Squir → Meshy/Tripo'],
        ['  Этап 5', 'Скрипт на bpy: импорт → bbox → сверка с эталоном → нормализация'],
        ['  Этап 6', 'Отчёт о покрытии: Σads_count готовых моделей / Σads_count'],
    ]
    for r in info_rows:
        ws_info.append(r)
    ws_info.column_dimensions['A'].width = 30
    ws_info.column_dimensions['B'].width = 90
    for r in range(1, ws_info.max_row + 1):
        ws_info.cell(row=r, column=1).font = Font(bold=True, color='1F4E79')
        for c in (1, 2):
            ws_info.cell(row=r, column=c).alignment = Alignment(vertical='top', wrap_text=True)

    # ---- Лист "MVP top-N" ----
    ws = wb.create_sheet(f'MVP top-{len(out)}')

    # 1-я строка — заголовки групп
    title_font = Font(bold=True, color='1F4E79', size=11)
    head_font = Font(bold=True, color='FFFFFF', size=10)
    head_fill = PatternFill('solid', fgColor='1F4E79')
    cidx = 1
    for group_name, color, cols in GROUPS:
        n = len(cols)
        ws.cell(row=1, column=cidx, value=group_name)
        if n > 1:
            ws.merge_cells(start_row=1, end_row=1,
                           start_column=cidx, end_column=cidx + n - 1)
        cell = ws.cell(row=1, column=cidx)
        cell.font = title_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.fill = PatternFill('solid', fgColor=color)
        for c in range(cidx, cidx + n):
            ws.cell(row=1, column=c).border = border
        cidx += n

    # 2-я строка — названия колонок
    col_order = [c for _, _, cols in GROUPS for c in cols]
    for c_idx, name in enumerate(col_order, start=1):
        cell = ws.cell(row=2, column=c_idx, value=name)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    ws.row_dimensions[2].height = 36

    # 3+ — данные
    df_out = out[col_order].copy()
    for r_idx, row in df_out.iterrows():
        for c_idx, val in enumerate(row, start=1):
            if pd.isna(val):
                v = None
            elif isinstance(val, (bool, np.bool_)):
                v = bool(val)
            elif isinstance(val, (np.integer,)):
                v = int(val)
            elif isinstance(val, (np.floating,)):
                v = float(val)
            else:
                v = val
            ws.cell(row=r_idx + 3, column=c_idx, value=v)

    # форматирование колонок
    money_cols = {'Средняя цена, ₽ (с инфл.)'}
    int_cols = {'Ранг', 'Поколение', 'Год от', 'Год до', 'Вариантов кузовов',
                'Кол-во объявлений', 'Длина, мм', 'Ширина, мм', 'Высота, мм',
                'Колёсная база, мм', 'Клиренс, мм'}
    score_cols = {'Б: локализация', 'Б: тип ТС', 'Б: объявления', 'Б: цена', 'Б: возраст'}
    score3_cols = {'Балл приоритета'}
    pct_cols = {'Покрытие, %'}
    age_cols = {'Средний возраст, лет'}

    left_align = {'ID', 'Марка', 'Модель', 'Рестайлинги', 'Типы кузова',
                  'Локализация', 'Тип ТС', 'Источник габаритов',
                  'Статус 3D', 'Источник 3D', 'Лицензия 3D', 'Путь к файлу',
                  'Формат', 'Проверка масштаба', 'Комментарий'}
    center_align = {'Реальные метрики', 'Совместимо с Blender'}

    last_row = ws.max_row
    last_col = ws.max_column
    for c_idx, name in enumerate(col_order, start=1):
        if name in money_cols:
            fmt = '#,##0\\ ₽'
        elif name in int_cols:
            fmt = '#,##0'
        elif name in score_cols:
            fmt = '0.000'
        elif name in score3_cols:
            fmt = '0.000'
        elif name in pct_cols:
            fmt = '0.00\\ %'
        elif name in age_cols:
            fmt = '0.0'
        else:
            fmt = None
        for r in range(3, last_row + 1):
            cell = ws.cell(row=r, column=c_idx)
            cell.border = border
            if fmt:
                cell.number_format = fmt
            if name in left_align:
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=False)
            elif name in center_align:
                cell.alignment = Alignment(horizontal='center', vertical='top')
            else:
                cell.alignment = Alignment(horizontal='right', vertical='top')

    # цветовые шкалы
    pcol = col_order.index('Балл приоритета') + 1
    ws.conditional_formatting.add(
        f'{get_column_letter(pcol)}3:{get_column_letter(pcol)}{last_row}',
        ColorScaleRule(start_type='min', start_color='F8696B',
                       mid_type='percentile', mid_value=50, mid_color='FFEB84',
                       end_type='max', end_color='63BE7B'),
    )
    ccol = col_order.index('Покрытие, %') + 1
    ws.conditional_formatting.add(
        f'{get_column_letter(ccol)}3:{get_column_letter(ccol)}{last_row}',
        ColorScaleRule(start_type='min', start_color='FFFFFF',
                       end_type='max', end_color='B4C7E7'),
    )

    # ширины
    widths = {
        'Ранг': 6, 'Балл приоритета': 12, 'Покрытие, %': 11, 'ID': 32,
        'Марка': 14, 'Модель': 18, 'Поколение': 9, 'Рестайлинги': 11,
        'Типы кузова': 38, 'Год от': 7, 'Год до': 7, 'Вариантов кузовов': 9,
        'Локализация': 14, 'Тип ТС': 14,
        'Кол-во объявлений': 13, 'Средняя цена, ₽ (с инфл.)': 19,
        'Средний возраст, лет': 13, 'Реальные метрики': 11,
        'Б: локализация': 10, 'Б: тип ТС': 10, 'Б: объявления': 11,
        'Б: цена': 10, 'Б: возраст': 10,
        'Длина, мм': 9, 'Ширина, мм': 9, 'Высота, мм': 9,
        'Колёсная база, мм': 12, 'Клиренс, мм': 10,
        'Источник габаритов': 18,
        'Статус 3D': 12, 'Источник 3D': 14, 'Лицензия 3D': 14,
        'Путь к файлу': 22, 'Формат': 9,
        'Проверка масштаба': 14, 'Совместимо с Blender': 12,
        'Комментарий': 30,
    }
    for c_idx, name in enumerate(col_order, start=1):
        ws.column_dimensions[get_column_letter(c_idx)].width = widths.get(name, 14)

    ws.auto_filter.ref = f'A2:{get_column_letter(last_col)}{last_row}'
    ws.freeze_panes = 'E3'  # фиксируем приоритет (4 первые колонки)

    # ---- Лист "Метрики покрытия" ----
    ws_m = wb.create_sheet('Метрики')
    by_origin = top.groupby('origin_category')['ads_count'].agg(['count', 'sum']).reset_index()
    by_origin['cov_%'] = by_origin['sum'] / total_ads * 100
    by_type = top.groupby('vehicle_type_category')['ads_count'].agg(['count', 'sum']).reset_index()
    by_type['cov_%'] = by_type['sum'] / total_ads * 100

    ws_m.append(['Срез', 'Категория', 'Авто в MVP', 'Объявлений', 'Покрытие %'])
    for cell in ws_m[1]:
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = border
    for _, r in by_origin.iterrows():
        ws_m.append(['Локализация', r['origin_category'], int(r['count']),
                     int(r['sum']), float(r['cov_%'])])
    for _, r in by_type.iterrows():
        ws_m.append(['Тип ТС', r['vehicle_type_category'], int(r['count']),
                     int(r['sum']), float(r['cov_%'])])
    ws_m.append(['Всего', '', len(top), int(top['ads_count'].sum()),
                 float(top['coverage_%'].iloc[-1])])
    for r in range(2, ws_m.max_row + 1):
        for c in range(1, 6):
            cell = ws_m.cell(row=r, column=c)
            cell.border = border
            if c == 5:
                cell.number_format = '0.00\\ %'
                cell.alignment = Alignment(horizontal='right')
            elif c in (3, 4):
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal='right')
    ws_m.column_dimensions['A'].width = 14
    ws_m.column_dimensions['B'].width = 18
    ws_m.column_dimensions['C'].width = 12
    ws_m.column_dimensions['D'].width = 14
    ws_m.column_dimensions['E'].width = 12

    wb.save(OUT)
    print(f'\n✅ Сохранено: {OUT}')
    print(f'   {len(out):,} строк × {len(col_order)} колонок')
    print(f'   Покрытие: {out["Покрытие, %"].iloc[-1]:.2f}%')
    print(f'   Размер: {OUT.stat().st_size/1024:.1f} КБ')


if __name__ == '__main__':
    main()
